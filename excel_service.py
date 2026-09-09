"""
Document Reference: excel_service.py
System: Project Chronicle - Enterprise AI Workflow Manager
Branch: A.14.12
Description: Generates multi-tab GxP-compliant Excel workbooks with audit logs.
"""

from io import BytesIO
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from chronicle_db import chronicle_db
from chronicle_models import AuditLogEntry, Task


def generate_portfolio_excel(tenant_id: str) -> bytes:
    tasks: List[Task] = chronicle_db.list_tasks(tenant_id)
    audit_logs: List[AuditLogEntry] = chronicle_db.get_audit_logs(tenant_id)
    summary = chronicle_db.get_portfolio_summary(tenant_id)

    wb = openpyxl.Workbook()

    # Style presets
    header_fill = PatternFill(
        start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    zebra_fill = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )

    # ==========================================
    # SHEET 1: PORTFOLIO SUMMARY
    # ==========================================
    ws1 = wb.active
    ws1.title = "Portfolio Overview"
    ws1.views.sheetView[0].showGridLines = True

    ws1.cell(row=1, column=1, value="PROJECT CHRONICLE - PORTFOLIO SUMMARY")
    ws1.cell(row=1, column=1).font = title_font
    ws1.cell(row=2, column=1, value=f"Branch: A.14.12 | Tenant: {tenant_id}")
    ws1.cell(row=2, column=1).font = Font(
        name="Calibri", size=10, italic=True, color="64748B"
    )

    metrics = [
        ("Total Managed Tasks", summary.total_tasks),
        ("Completion Rate (%)", f"{summary.completion_rate_percent}%"),
        ("Pending Human Approvals (HITL)", summary.pending_approvals),
        ("Active Milestones", len(summary.milestones)),
    ]

    ws1.cell(row=4, column=1, value="Executive KPI Metrics").font = Font(
        name="Calibri", size=11, bold=True
    )
    ws1.cell(row=4, column=1).fill = PatternFill(
        start_color="E2E8F0", end_color="E2E8F0", fill_type="solid"
    )
    ws1.cell(row=4, column=2, value="Value").font = Font(
        name="Calibri", size=11, bold=True
    )
    ws1.cell(row=4, column=2).fill = PatternFill(
        start_color="E2E8F0", end_color="E2E8F0", fill_type="solid"
    )

    for idx, (label, val) in enumerate(metrics, start=5):
        ws1.cell(row=idx, column=1, value=label).font = regular_font
        ws1.cell(row=idx, column=1).border = thin_border
        ws1.cell(row=idx, column=2, value=val).font = Font(
            name="Calibri", size=10, bold=True
        )
        ws1.cell(row=idx, column=2).border = thin_border

    # Status distribution
    ws1.cell(row=10, column=1, value="Status Breakdown").font = Font(
        name="Calibri", size=11, bold=True
    )
    ws1.cell(row=10, column=1).fill = PatternFill(
        start_color="E2E8F0", end_color="E2E8F0", fill_type="solid"
    )
    ws1.cell(row=10, column=2, value="Count").font = Font(
        name="Calibri", size=11, bold=True
    )
    ws1.cell(row=10, column=2).fill = PatternFill(
        start_color="E2E8F0", end_color="E2E8F0", fill_type="solid"
    )

    for idx, (st, cnt) in enumerate(summary.by_status.items(), start=11):
        ws1.cell(row=idx, column=1, value=st).font = regular_font
        ws1.cell(row=idx, column=1).border = thin_border
        ws1.cell(row=idx, column=2, value=cnt).font = regular_font
        ws1.cell(row=idx, column=2).border = thin_border

    # ==========================================
    # SHEET 2: KANBAN & TASK SCHEDULE
    # ==========================================
    ws2 = wb.create_sheet(title="Kanban Tasks")
    ws2.views.sheetView[0].showGridLines = True

    task_headers = [
        "Task ID",
        "Title",
        "Category",
        "Status",
        "Priority",
        "Milestone",
        "Assignee",
        "Due Date",
        "Approval Status",
        "Reviewer Comments",
    ]

    for col_idx, h in enumerate(task_headers, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, t in enumerate(tasks, start=2):
        row_values = [
            t.id[:8] + "...",
            t.title,
            t.category.value,
            t.status.value,
            t.priority.value,
            t.milestone,
            t.assignee,
            t.due_date or "N/A",
            t.approval_status,
            t.reviewer_comments or "",
        ]
        use_zebra = row_idx % 2 == 0
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if use_zebra:
                cell.fill = zebra_fill

    # ==========================================
    # SHEET 3: 21 CFR PART 11 GXP AUDIT TRAIL
    # ==========================================
    ws3 = wb.create_sheet(title="GxP Audit Trail")
    ws3.views.sheetView[0].showGridLines = True

    audit_headers = [
        "Audit Log ID",
        "Timestamp (UTC)",
        "Actor",
        "Action",
        "Target Task ID",
        "Comments",
        "Previous State Snapshot",
        "New State Snapshot",
    ]

    for col_idx, h in enumerate(audit_headers, start=1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, a in enumerate(audit_logs, start=2):
        row_values = [
            a.id[:8] + "...",
            a.timestamp,
            a.actor,
            a.action,
            a.task_id[:8] + "...",
            a.comments,
            a.previous_state or "",
            a.new_state or "",
        ]
        use_zebra = row_idx % 2 == 0
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if use_zebra:
                cell.fill = zebra_fill

    # Auto-fit columns across all sheets
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
