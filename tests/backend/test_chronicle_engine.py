"""
Document Reference: tests/backend/test_chronicle_engine.py
System: Project Chronicle - Enterprise AI Workflow Manager
Branch: A.14.12
Description: Validation suite covering Task CRUD, Circuit Breakers (RM-005),
             GxP Audit Logs, Excel Reporting, and REST API controllers.
"""

from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

from chronicle_db import chronicle_db
from chronicle_models import (
    TaskCategory,
    TaskCreate,
    TaskPriority,
    TaskStatus,
    TaskUpdate,
)
from circuit_breaker import CircuitState, TaskServiceCircuitBreaker
from excel_service import generate_portfolio_excel
from server import server

client = TestClient(server)
TEST_TENANT = "55555555-5555-5555-5555-555555555555"


# ==========================================
# 1. DATABASE & AUDIT TRAIL TESTS
# ==========================================
def test_task_creation_and_audit_logging() -> None:
    """Verifies task creation generates proper 21 CFR Part 11 audit records."""
    task_in = TaskCreate(
        title="Draft IQ Protocol for AI Host Server",
        description="Verify hardware, Python dependencies, and networking.",
        category=TaskCategory.ACTION,
        priority=TaskPriority.HIGH,
        milestone="Sprint 1 - Foundation",
        assignee="Dave Barbour",
        due_date="2026-09-30",
        tenant_id=TEST_TENANT,
    )
    task = chronicle_db.create_task(task_in, actor="Lead Tester")
    assert task.title == task_in.title
    assert task.category == TaskCategory.ACTION
    assert task.status == TaskStatus.IN_REVIEW  # Actions require review

    # Verify audit log
    logs = chronicle_db.get_audit_logs(TEST_TENANT, task_id=task.id)
    assert len(logs) >= 1
    assert logs[0].action == "CREATE"
    assert logs[0].actor == "Lead Tester"


def test_task_status_transition_and_audit() -> None:
    """Verifies task updates generate delta audit trail entries."""
    task_in = TaskCreate(
        title="OQ Execution: API Boundary Validation",
        category=TaskCategory.SCHEDULE,
        priority=TaskPriority.MEDIUM,
        tenant_id=TEST_TENANT,
    )
    task = chronicle_db.create_task(task_in, actor="Dave Barbour")

    # Update status to IN_PROGRESS
    updated = chronicle_db.update_task(
        task.id,
        TEST_TENANT,
        TaskUpdate(status=TaskStatus.IN_PROGRESS),
        actor="Dave Barbour",
    )
    assert updated is not None
    assert updated.status == TaskStatus.IN_PROGRESS

    # Verify audit log recorded update
    logs = chronicle_db.get_audit_logs(TEST_TENANT, task_id=task.id)
    assert len(logs) >= 2
    assert logs[0].action == "UPDATE"


def test_task_deletion() -> None:
    """Verifies task deletion and corresponding DELETE audit record."""
    task_in = TaskCreate(
        title="Temporary Spike Task",
        category=TaskCategory.RESEARCH,
        tenant_id=TEST_TENANT,
    )
    task = chronicle_db.create_task(task_in)
    deleted = chronicle_db.delete_task(task.id, TEST_TENANT, actor="Admin")
    assert deleted is True

    # Confirm it cannot be retrieved
    assert chronicle_db.get_task(task.id, TEST_TENANT) is None

    # Verify deletion audit log
    logs = chronicle_db.get_audit_logs(TEST_TENANT, task_id=task.id)
    assert logs[0].action == "DELETE"


# ==========================================
# 2. CIRCUIT BREAKER RESILIENCE (RM-005)
# ==========================================
def test_circuit_breaker_trip_and_recovery() -> None:
    """Verifies circuit breaker trips on consecutive errors and blocks calls."""
    breaker = TaskServiceCircuitBreaker(
        name="TestBreaker",
        failure_threshold=2,
        recovery_timeout_seconds=0.1,
    )

    def failing_call() -> None:
        raise ConnectionError("Simulated integration failure")

    # 1st failure
    with pytest.raises(ConnectionError):
        breaker.call(failing_call)
    assert breaker.state == CircuitState.CLOSED
    assert breaker.failure_count == 1

    # 2nd failure -> should trip to OPEN
    with pytest.raises(ConnectionError):
        breaker.call(failing_call)
    assert breaker.state == CircuitState.OPEN

    # 3rd call immediately -> rejected by circuit breaker with RuntimeError
    with pytest.raises(RuntimeError) as exc_info:
        breaker.call(failing_call)
    assert "is OPEN" in str(exc_info.value)

    # Wait for recovery timeout
    import time

    time.sleep(0.15)

    # Trial probe success
    result = breaker.call(lambda: "Recovered")
    assert result == "Recovered"
    assert breaker.state == CircuitState.CLOSED


# ==========================================
# 3. EXCEL PORTFOLIO EXPORT
# ==========================================
def test_excel_export_structure() -> None:
    """Verifies multi-tab workbook generation with data integrity."""
    excel_bytes = generate_portfolio_excel(TEST_TENANT)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Portfolio Overview" in sheet_names
    assert "Kanban Tasks" in sheet_names
    assert "GxP Audit Trail" in sheet_names

    ws1 = wb["Portfolio Overview"]
    assert "PROJECT CHRONICLE" in str(ws1.cell(row=1, column=1).value)


# ==========================================
# 4. REST API ENDPOINTS
# ==========================================
def test_api_tasks_crud_lifecycle() -> None:
    """Verifies complete REST API CRUD operations for tasks."""
    # List tasks (triggers auto-seeding)
    list_resp = client.get(f"/api/tasks?tenant_id={TEST_TENANT}")
    assert list_resp.status_code == 200
    tasks = list_resp.json()
    assert len(tasks) >= 1

    # Create task via API
    new_task = {
        "title": "API Integration Test Task",
        "description": "Created via FastAPI endpoint.",
        "category": "ACTION",
        "priority": "HIGH",
        "milestone": "Sprint 1 - Foundation",
        "assignee": "Dave Barbour",
        "tenant_id": TEST_TENANT,
    }
    create_resp = client.post("/api/tasks", json=new_task)
    assert create_resp.status_code == 200
    task_id = create_resp.json()["id"]

    # Patch task (Move to IN_PROGRESS)
    patch_resp = client.patch(
        f"/api/tasks/{task_id}?tenant_id={TEST_TENANT}",
        json={"status": "IN_PROGRESS"},
    )
    assert patch_resp.status_code == 200
    assert patch_resp.json()["status"] == "IN_PROGRESS"

    # Get Portfolio Summary
    summary_resp = client.get(f"/api/portfolio/summary?tenant_id={TEST_TENANT}")
    assert summary_resp.status_code == 200
    summary = summary_resp.json()
    assert summary["total_tasks"] >= 1

    # Download Excel Report
    excel_resp = client.get(f"/api/tasks/export/excel?tenant_id={TEST_TENANT}")
    assert excel_resp.status_code == 200
    assert "application/vnd.openxmlformats" in excel_resp.headers["content-type"]

    # Health check
    health_resp = client.get("/api/health")
    assert health_resp.status_code == 200
    assert health_resp.json()["status"] == "UP"
